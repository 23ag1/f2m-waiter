from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from f2m_engine.config.taxonomy import RuntimeTaxonomy


@dataclass(frozen=True)
class DishFeaturesBuildStats:
    parsed_rows: int
    dishes_processed: int
    skipped_rows: int
    skipped_by_reason: dict[str, int]
    unique_dish_ids: int
    nutrition_parse_success_rate: float
    feature_rows_written: int
    cache_rows_written: int


ALIASES: dict[str, str] = {
    "томат": "tomato",
    "томаты": "tomato",
    "томаты черри": "tomato",
    "черри": "tomato",
    "помидор": "tomato",
    "сыр": "cheese",
    "моцарелла": "cheese",
    "пармезан": "cheese",
    "сулугуни": "cheese",
    "лосось": "salmon",
    "семга": "salmon",
    "креветки": "shrimp",
    "креветка": "shrimp",
    "тунец": "tuna",
    "говядина": "beef",
    "баранина": "lamb",
    "свинина": "pork",
    "курица": "chicken",
    "куриный": "chicken",
    "индейка": "turkey",
    "бекон": "bacon",
    "кимчи": "kimchi",
    "кунжут": "sesame",
    "семена кунжута": "sesame",
    "нут": "chickpea",
    "чечевица": "lentil",
    "яйцо": "egg",
    "яйца": "egg",
    "молоко": "milk",
    "сливки": "cream",
    "масло": "butter",
    "сахар": "sugar",
    "мука": "wheat",
    "пшеничное тесто": "wheat",
    "рис": "rice",
    "лапша": "noodles",
    "паста": "pasta",
    "брокколи": "broccoli",
    "лук": "onion",
    "грибы": "mushroom",
    "шампиньоны": "mushroom",
    "чеснок": "garlic",
    "огурец": "cucumber",
    "зелень": "greens",
    "соевый соус": "soy_sauce",
    "соус терияки": "teriyaki",
    "терияки": "teriyaki",
    "устричный соус": "oyster_sauce",
    "рыбный соус": "fish_sauce",
    "кокосовое молоко": "coconut_milk",
    "чили": "chili",
    "карри": "curry",
}


HARD_FLAG_RULES: dict[str, set[str]] = {
    "gluten": {"wheat"},
    "lactose": {"milk", "cream", "cheese", "butter"},
    "peanut": {"peanut"},
    "nuts": {"nuts"},
    "egg": {"egg"},
    "soy": {"soy_sauce", "teriyaki"},
    "fish_seafood": {"salmon", "tuna", "shrimp", "fish_sauce", "oyster_sauce"},
    "pork": {"pork"},
}


ASIAN_EXPLICIT_KEYWORDS = (
    "вок",
    "кимчи",
    "терияки",
    "соевый соус",
    "том-ям",
    "том ям",
    "рамен",
    "бань бао",
    "чапче",
)

CUISINE_KEYWORDS: dict[str, tuple[str, ...]] = {
    "asian": ASIAN_EXPLICIT_KEYWORDS,
    "italian": ("пицца", "карбонара", "пенне", "спагетти", "тирамису", "помодоро"),
    "georgian_caucasian": ("хачапури", "хинкали", "лобио", "люля", "кебаб", "кавказ"),
    "russian_home": ("борщ", "пельмени", "сырники"),
    "fast_food": ("бургер", "фри", "стритфуд"),
    "street_food": ("стритфуд", "шаверма", "шаурма"),
    "home_style": ("домашн", "традицион"),
}


FORMAT_KEYWORDS: dict[str, tuple[str, ...]] = {
    "soup": ("суп", "том-ям", "борщ", "рамен"),
    "salad": ("салат",),
    "pizza": ("пицца",),
    "burger": ("бургер",),
    "roll": ("ролл",),
    "baked": ("запеч", "печен", "духов", "пицца"),
    "fried": ("жарен", "фритюр"),
    "grilled": ("гриль", "шашлык", "кебаб"),
    "creamy": ("сливоч", "крем", "сыр"),
    "hearty": ("плотн", "сытн", "наварист"),
    "light": ("легк",),
    "dessert": ("тирамису", "десерт", "торт"),
    "breakfast": ("сырники", "завтрак"),
}

PRIMARY_INGREDIENT_FROM_NAME: dict[str, str] = {
    "говядин": "beef",
    "куриц": "chicken",
    "бекон": "bacon",
    "кревет": "shrimp",
    "лосос": "salmon",
    "свинин": "pork",
    "баранин": "lamb",
    "тунц": "tuna",
    "индейк": "turkey",
}


def build_dish_features_stage_a(
    menu_path: Path | str,
    output_dir: Path | str,
    taxonomy: RuntimeTaxonomy,
) -> DishFeaturesBuildStats:
    loaded = _load_menu_xlsx(Path(menu_path))
    dishes = loaded["dishes"]
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    audit_dir = out / "audit"
    audit_dir.mkdir(parents=True, exist_ok=True)
    now = datetime.now(timezone.utc).isoformat()

    feature_rows: list[dict[str, Any]] = []
    cache_rows: list[dict[str, Any]] = []
    ingredient_seen: set[str] = set()
    evidence_rows: list[list[Any]] = []
    missing_primary_rows: list[list[Any]] = []

    for dish in dishes:
        dish_id = str(dish["dish_id"])
        name = str(dish.get("dish_name", ""))
        description = str(dish.get("description", ""))
        category = str(dish.get("category", ""))
        ingredients_text = str(dish.get("ingredients_text", ""))
        normalized_name = _normalize(name)
        normalized_category = _normalize(category)
        normalized_description = _normalize(description)
        normalized_ingredients = _normalize(ingredients_text)
        concat = " ".join(
            value for value in [normalized_name, normalized_description, normalized_category, normalized_ingredients] if value
        )

        ingredient_weights, aliases_used = _extract_ingredients(normalized_ingredients)
        _ensure_primary_ingredient_from_name(
            name_text=normalized_name,
            ingredient_weights=ingredient_weights,
            missing_primary_rows=missing_primary_rows,
            dish_id=dish_id,
            dish_name=name,
            evidence_rows=evidence_rows,
        )
        ingredient_seen.update(ingredient_weights.keys())

        nutrition = _compute_nutrition(dish=dish)
        cuisine_weights = _extract_cuisine(
            name=normalized_name,
            description=normalized_description,
            category=normalized_category,
            ingredients=normalized_ingredients,
            nutrition=nutrition,
            taxonomy=taxonomy,
            dish_id=dish_id,
            dish_name=name,
            evidence_rows=evidence_rows,
        )
        format_weights = _extract_format_texture(
            name=normalized_name,
            category=normalized_category,
            ingredients=normalized_ingredients,
            taxonomy=taxonomy,
            dish_id=dish_id,
            dish_name=name,
            evidence_rows=evidence_rows,
        )
        taste_weights = _extract_taste_from_ingredients(ingredient_weights)
        context_weights = _extract_context_from_text(concat)
        hard_flags = _compute_hard_flags(ingredient_weights=ingredient_weights, concat=concat, taxonomy=taxonomy)

        feature_rows.extend(
            _to_feature_rows(
                dish_id=dish_id,
                axis="taste",
                values=taste_weights,
                now=now,
                feature_version=taxonomy.version,
            )
        )
        feature_rows.extend(
            _to_feature_rows(
                dish_id=dish_id,
                axis="ingredient",
                values=ingredient_weights,
                now=now,
                feature_version=taxonomy.version,
            )
        )
        feature_rows.extend(
            _to_feature_rows(
                dish_id=dish_id,
                axis="cuisine",
                values=cuisine_weights,
                now=now,
                feature_version=taxonomy.version,
            )
        )
        feature_rows.extend(
            _to_feature_rows(
                dish_id=dish_id,
                axis="format_texture",
                values=format_weights,
                now=now,
                feature_version=taxonomy.version,
            )
        )
        feature_rows.extend(
            _to_feature_rows(
                dish_id=dish_id,
                axis="context",
                values=context_weights,
                now=now,
                feature_version=taxonomy.version,
            )
        )
        for ingredient_key, weight in ingredient_weights.items():
            evidence_rows.append(
                [dish_id, name, "ingredient", ingredient_key, "ingredient_text_or_name", weight, "rule"]
            )

        for key, value in hard_flags.items():
            feature_rows.append(
                {
                    "dish_id": dish_id,
                    "axis": "restriction",
                    "feature_key": key,
                    "value_bool": bool(value),
                    "value_num": None,
                    "source": "rule",
                    "confidence": 1.0,
                    "feature_version": taxonomy.version,
                    "updated_at": now,
                }
            )
            evidence_rows.append([dish_id, name, "restriction", key, "ingredient_rule", float(value), "rule"])

        for key in ("low_calorie", "high_protein", "low_fat", "high_carb"):
            tag_value = nutrition.get(key)
            if tag_value is None:
                continue
            feature_rows.append(
                {
                    "dish_id": dish_id,
                    "axis": "nutrition",
                    "feature_key": key,
                    "value_bool": bool(tag_value),
                    "value_num": None,
                    "source": "computed",
                    "confidence": 1.0,
                    "feature_version": taxonomy.version,
                    "updated_at": now,
                }
            )
            evidence_rows.append([dish_id, name, "nutrition", key, "deterministic_formula", float(tag_value), "computed"])
        if nutrition.get("satiety_index") is not None:
            feature_rows.append(
                {
                    "dish_id": dish_id,
                    "axis": "nutrition",
                    "feature_key": "satiety_index",
                    "value_num": float(nutrition["satiety_index"]),
                    "value_bool": None,
                    "source": "computed",
                    "confidence": 1.0,
                    "feature_version": taxonomy.version,
                    "updated_at": now,
                }
            )
            evidence_rows.append(
                [dish_id, name, "nutrition", "satiety_index", "deterministic_formula", float(nutrition["satiety_index"]), "computed"]
            )

        cache_rows.append(
            {
                "dish_id": dish_id,
                "taste": taste_weights,
                "ingredient": ingredient_weights,
                "cuisine": cuisine_weights,
                "format_texture": format_weights,
                "context": context_weights,
                "hard_flags": hard_flags,
                "nutrition": nutrition,
                "meta": {
                    "name": name,
                    "restaurant": dish.get("restaurant"),
                    "category": category,
                    "feature_version": taxonomy.version,
                    "aliases_used": aliases_used,
                },
            }
        )

    _write_jsonl(out / "dish_feature_values.jsonl", feature_rows)
    (out / "dish_features_cache.json").write_text(
        json.dumps(cache_rows, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (out / "ingredient_dictionary.json").write_text(
        json.dumps(sorted(ingredient_seen), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    (out / "ingredient_aliases.json").write_text(
        json.dumps(ALIASES, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    _write_dish_audit(
        audit_dir=audit_dir,
        loaded=loaded,
        dishes=dishes,
        cache_rows=cache_rows,
        evidence_rows=evidence_rows,
        missing_primary_rows=missing_primary_rows,
    )
    _write_false_positive_watchlist(audit_dir=audit_dir, cache_rows=cache_rows)

    nutrition_rate = 0.0
    if len(dishes) > 0:
        nutrition_rate = loaded["nutrition_success"] / len(dishes)
    return DishFeaturesBuildStats(
        parsed_rows=loaded["parsed_rows"],
        dishes_processed=len(dishes),
        skipped_rows=loaded["skipped_rows"],
        skipped_by_reason=loaded["skipped_by_reason"],
        unique_dish_ids=len({str(d["dish_id"]) for d in dishes}),
        nutrition_parse_success_rate=round(nutrition_rate, 4),
        feature_rows_written=len(feature_rows),
        cache_rows_written=len(cache_rows),
    )


def _load_menu_xlsx(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Меню" not in workbook.sheetnames:
        raise ValueError("Required sheet 'Меню' not found")
    sheet = workbook["Меню"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return {
            "dishes": [],
            "parsed_rows": 0,
            "skipped_rows": 0,
            "skipped_by_reason": {},
            "skipped_rows_detail": [],
            "nutrition_parse_errors": [],
            "nutrition_success": 0,
            "ingredient_alias_review": [],
        }

    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    index = {header: idx for idx, header in enumerate(headers)}

    def pick(row: tuple[Any, ...], *names: str) -> Any:
        for name in names:
            idx = index.get(name)
            if idx is not None and idx < len(row):
                return row[idx]
        return None

    result: list[dict[str, Any]] = []
    seen_dish_ids: set[str] = set()
    skipped_rows_detail: list[dict[str, Any]] = []
    skipped_by_reason: dict[str, int] = {}
    nutrition_parse_errors: list[dict[str, Any]] = []
    nutrition_success = 0
    review_counter: dict[str, dict[str, Any]] = {}

    def skip(reason: str, row_idx: int, row: tuple[Any, ...], dish_id: Any) -> None:
        skipped_by_reason[reason] = skipped_by_reason.get(reason, 0) + 1
        skipped_rows_detail.append(
            {
                "row_index": row_idx,
                "reason": reason,
                "dish_id": dish_id,
                "dish_name": pick(row, "Название блюда", "Блюдо", "Наименование"),
            }
        )

    for row_idx, row in enumerate(rows[1:], start=2):
        dish_id = pick(row, "ID блюда")
        if dish_id in (None, ""):
            skip("missing_dish_id", row_idx, row, dish_id)
            continue
        dish_id = str(dish_id).strip()
        if dish_id in seen_dish_ids:
            skip("duplicate_dish_id", row_idx, row, dish_id)
            continue
        seen_dish_ids.add(dish_id)

        dish_name = pick(row, "Название блюда", "Название", "Блюдо", "Наименование")
        if dish_name in (None, ""):
            skip("missing_name", row_idx, row, dish_id)
            continue

        ingredients = pick(row, "Ингредиенты", "Состав", "Описание состава")
        if ingredients in (None, ""):
            skip("missing_ingredients", row_idx, row, dish_id)
            continue

        kcal_raw = pick(row, "Ккал")
        nutrition = _parse_nutrition_from_kcal(kcal_raw)
        if nutrition["ok"]:
            nutrition_success += 1
        else:
            nutrition_parse_errors.append(
                {
                    "row_index": row_idx,
                    "dish_id": dish_id,
                    "dish_name": str(dish_name),
                    "raw_kcal": kcal_raw,
                }
            )
        _collect_alias_review(str(ingredients), review_counter)
        weight_raw = pick(row, "Вес")
        weight_parse = _parse_total_weight_grams(weight_raw)
        if weight_parse["needs_review"]:
            nutrition_parse_errors.append(
                {
                    "row_index": row_idx,
                    "dish_id": dish_id,
                    "dish_name": str(dish_name),
                    "raw_kcal": f"weight_parse:{weight_raw}",
                }
            )

        result.append(
            {
                "dish_id": dish_id,
                "dish_name": dish_name,
                "restaurant": pick(row, "Ресторан"),
                "description": pick(row, "Описание", "Описание блюда"),
                "category": pick(row, "Категория"),
                "ingredients_text": ingredients,
                "kcal_per_portion": nutrition["kcal"],
                "protein_per_portion": nutrition["protein"],
                "fat_per_portion": nutrition["fat"],
                "carb_per_portion": nutrition["carb"],
                "protein_per_100g": _per_100(nutrition["protein"], weight_parse["grams"]),
                "fat_per_100g": _per_100(nutrition["fat"], weight_parse["grams"]),
                "carb_per_100g": _per_100(nutrition["carb"], weight_parse["grams"]),
                "weight_g": weight_parse["grams"],
                "weight_needs_review": weight_parse["needs_review"],
                "weight_raw": weight_raw,
            }
        )
    return {
        "dishes": result,
        "parsed_rows": max(len(rows) - 1, 0),
        "skipped_rows": len(skipped_rows_detail),
        "skipped_by_reason": skipped_by_reason,
        "skipped_rows_detail": skipped_rows_detail,
        "nutrition_parse_errors": nutrition_parse_errors,
        "nutrition_success": nutrition_success,
        "ingredient_alias_review": list(review_counter.values()),
    }


def _extract_ingredients(concat_text: str) -> tuple[dict[str, float], list[str]]:
    normalized = re.sub(r"[^a-zа-я0-9\s-]", " ", concat_text)
    normalized = re.sub(r"\s+", " ", normalized).strip()

    weights: dict[str, float] = {}
    aliases_used: list[str] = []
    for alias, canonical in ALIASES.items():
        if alias in normalized:
            weights[canonical] = max(weights.get(canonical, 0.0), 1.0)
            aliases_used.append(alias)
    return weights, sorted(set(aliases_used))


def _ensure_primary_ingredient_from_name(
    name_text: str,
    ingredient_weights: dict[str, float],
    missing_primary_rows: list[list[Any]],
    dish_id: str,
    dish_name: str,
    evidence_rows: list[list[Any]],
) -> None:
    for marker, canonical in PRIMARY_INGREDIENT_FROM_NAME.items():
        if marker not in name_text:
            continue
        if canonical in ingredient_weights:
            continue
        ingredient_weights[canonical] = 0.6
        missing_primary_rows.append([dish_id, dish_name, canonical, "added_from_name"])
        evidence_rows.append([dish_id, dish_name, "ingredient", canonical, "name_primary", 0.6, "rule"])


def _normalize(value: Any) -> str:
    return str(value or "").lower().replace("ё", "е").strip()


def _parse_nutrition_from_kcal(raw_value: Any) -> dict[str, Any]:
    if raw_value in (None, ""):
        return {"ok": False, "carb": None, "protein": None, "fat": None, "kcal": None}
    text = str(raw_value).replace("ё", "е")
    pattern = re.compile(
        r"К:\s*([0-9]+(?:[.,][0-9]+)?)\s*г[;,]\s*Б:\s*([0-9]+(?:[.,][0-9]+)?)\s*г[;,]\s*Ж:\s*([0-9]+(?:[.,][0-9]+)?)\s*г[;,]\s*Э:\s*([0-9]+(?:[.,][0-9]+)?)\s*ккал",
        re.IGNORECASE,
    )
    match = pattern.search(text)
    if match:
        carb, protein, fat, kcal = match.groups()
        return {
            "ok": True,
            "carb": float(carb.replace(",", ".")),
            "protein": float(protein.replace(",", ".")),
            "fat": float(fat.replace(",", ".")),
            "kcal": float(kcal.replace(",", ".")),
        }

    # Fallback for malformed strings: parse first four numbers as K,B,Ж,Э order.
    numbers = re.findall(r"([0-9]+(?:[.,][0-9]+)?)", text)
    if len(numbers) >= 4:
        carb, protein, fat, kcal = numbers[0], numbers[1], numbers[2], numbers[3]
    else:
        return {"ok": False, "carb": None, "protein": None, "fat": None, "kcal": None}

    return {
        "ok": True,
        "carb": float(carb.replace(",", ".")),
        "protein": float(protein.replace(",", ".")),
        "fat": float(fat.replace(",", ".")),
        "kcal": float(kcal.replace(",", ".")),
    }


def _parse_total_weight_grams(raw_value: Any) -> dict[str, Any]:
    if raw_value in (None, ""):
        return {"grams": None, "needs_review": True}
    text = str(raw_value).strip().lower().replace("ё", "е")
    if re.search(r"\d\s*-\s*\d", text):
        return {"grams": None, "needs_review": True}
    numbers = [float(item.replace(",", ".")) for item in re.findall(r"([0-9]+(?:[.,][0-9]+)?)", text)]
    if not numbers:
        return {"grams": None, "needs_review": True}
    grams = sum(numbers)
    return {"grams": grams, "needs_review": False}


def _per_100(value_per_portion: float | None, weight_g: float | None) -> float | None:
    if value_per_portion is None or weight_g is None or weight_g <= 0:
        return None
    return round(value_per_portion * 100.0 / weight_g, 2)


def _extract_keyword_axis(
    concat_text: str,
    mapping: dict[str, tuple[str, ...]],
    taxonomy: RuntimeTaxonomy,
    axis: str,
) -> dict[str, float]:
    values: dict[str, float] = {}
    for tag, keywords in mapping.items():
        if not taxonomy.is_allowed(axis, tag):
            continue
        if any(keyword in concat_text for keyword in keywords):
            values[tag] = 1.0
    return values


def _extract_cuisine(
    name: str,
    description: str,
    category: str,
    ingredients: str,
    nutrition: dict[str, Any],
    taxonomy: RuntimeTaxonomy,
    dish_id: str,
    dish_name: str,
    evidence_rows: list[list[Any]],
) -> dict[str, float]:
    concat = f"{name} {description} {category} {ingredients}"
    values: dict[str, float] = {}
    for tag, keywords in CUISINE_KEYWORDS.items():
        if not taxonomy.is_allowed("cuisine", tag):
            continue
        hit = next((keyword for keyword in keywords if keyword in concat), None)
        if hit:
            values[tag] = 1.0
            evidence_rows.append([dish_id, dish_name, "cuisine", tag, f"keyword:{hit}", 1.0, "rule"])
    explicit_healthy_keywords = ("здоров", "полезн", "фитнес", "правильн", "диет", "на пару", "су-вид")
    has_healthy_keyword = any(keyword in f"{name} {description} {category}" for keyword in explicit_healthy_keywords)
    strict_nutrition_healthy = bool(
        nutrition.get("low_calorie") is True
        and nutrition.get("high_protein") is True
        and nutrition.get("low_fat") is True
    )
    if taxonomy.is_allowed("cuisine", "healthy") and (has_healthy_keyword or strict_nutrition_healthy):
        values["healthy"] = 1.0
        reason = "keyword:healthy" if has_healthy_keyword else "nutrition_conjunction"
        evidence_rows.append([dish_id, dish_name, "cuisine", "healthy", reason, 1.0, "rule"])
    return values


def _extract_format_texture(
    name: str,
    category: str,
    ingredients: str,
    taxonomy: RuntimeTaxonomy,
    dish_id: str,
    dish_name: str,
    evidence_rows: list[list[Any]],
) -> dict[str, float]:
    values: dict[str, float] = {}
    for tag, keywords in FORMAT_KEYWORDS.items():
        if not taxonomy.is_allowed("format_texture", tag):
            continue
        search_scope = f"{name} {category}"
        if tag != "salad":
            search_scope = f"{search_scope} {ingredients}"
        hit = next((keyword for keyword in keywords if keyword in search_scope), None)
        if hit:
            values[tag] = 1.0
            evidence_rows.append([dish_id, dish_name, "format_texture", tag, f"keyword:{hit}", 1.0, "rule"])
    return values


def _extract_taste_from_ingredients(ingredient_weights: dict[str, float]) -> dict[str, float]:
    taste = {"sweet": 0.0, "salty": 0.0, "sour": 0.0, "spicy": 0.0, "bitter": 0.0, "umami": 0.0}
    if "sugar" in ingredient_weights:
        taste["sweet"] = 0.8
    if {"soy_sauce", "fish_white", "shrimp", "tuna", "salmon"} & ingredient_weights.keys():
        taste["umami"] = max(taste["umami"], 0.7)
        taste["salty"] = max(taste["salty"], 0.5)
    if "chili" in ingredient_weights or "curry" in ingredient_weights:
        taste["spicy"] = 0.8
    if "tomato" in ingredient_weights:
        taste["sour"] = max(taste["sour"], 0.3)
    return {key: value for key, value in taste.items() if value > 0}


def _extract_context_from_text(concat_text: str) -> dict[str, float]:
    context_map = {
        "morning": ("завтрак", "утрен"),
        "lunch": ("обед", "ланч"),
        "evening": ("ужин", "вечер"),
        "quick": ("быстр",),
        "romantic": ("романтич",),
        "healthy_choice": ("полезн", "легк"),
    }
    values: dict[str, float] = {}
    for tag, keywords in context_map.items():
        if any(keyword in concat_text for keyword in keywords):
            values[tag] = 0.6
    return values


def _compute_hard_flags(
    ingredient_weights: dict[str, float],
    concat: str,
    taxonomy: RuntimeTaxonomy,
) -> dict[str, bool]:
    flags: dict[str, bool] = {}
    ingredients = set(ingredient_weights.keys())
    for key, trigger in HARD_FLAG_RULES.items():
        if not taxonomy.is_allowed("restriction", key):
            continue
        flags[key] = bool(ingredients & trigger)
    # Alcohol is explicit text-based in MVP Stage A.
    if taxonomy.is_allowed("restriction", "no_sugar"):
        # Keep restriction axis closed; this is not a hard flag but a convenience signal.
        pass
    if "вино" in concat or "алкогол" in concat:
        if taxonomy.is_allowed("restriction", "halal"):
            flags["halal"] = False
    return flags


def _compute_nutrition(dish: dict[str, Any]) -> dict[str, Any]:
    kcal = dish.get("kcal_per_portion")
    protein = dish.get("protein_per_100g")
    fat = dish.get("fat_per_100g")
    carb = dish.get("carb_per_100g")
    needs_review = bool(dish.get("weight_needs_review"))

    low_calorie = bool(kcal is not None and kcal <= 200)
    if needs_review or protein is None or fat is None or carb is None:
        high_protein = None
        low_fat = None
        high_carb = None
        satiety_index = None
        satiety_class = None
    else:
        high_protein = protein >= 15
        low_fat = fat <= 5
        high_carb = carb >= 25 and fat <= 5
        satiety_index = round(float((protein * 3) + (carb * 0.4) - (fat * 0.3)), 2)
        if satiety_index <= 20:
            satiety_class = "snack"
        elif satiety_index <= 50:
            satiety_class = "light"
        else:
            satiety_class = "hearty"
    if satiety_index is None:
        satiety_class = None

    return {
        "weight_g": dish.get("weight_g"),
        "weight_needs_review": needs_review,
        "weight_raw": dish.get("weight_raw"),
        "kcal_per_portion": kcal,
        "protein_per_100g": protein,
        "fat_per_100g": fat,
        "carb_per_100g": carb,
        "low_calorie": low_calorie,
        "high_protein": high_protein,
        "low_fat": low_fat,
        "high_carb": high_carb,
        "satiety_index": satiety_index,
        "satiety_class": satiety_class,
    }


def _to_feature_rows(
    dish_id: str,
    axis: str,
    values: dict[str, float],
    now: str,
    feature_version: str,
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for key, weight in values.items():
        rows.append(
            {
                "dish_id": dish_id,
                "axis": axis,
                "feature_key": key,
                "value_num": float(weight),
                "value_bool": None,
                "source": "rule",
                "confidence": 1.0,
                "feature_version": feature_version,
                "updated_at": now,
            }
        )
    return rows


def _write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8") as file_obj:
        for row in rows:
            file_obj.write(json.dumps(row, ensure_ascii=False) + "\n")


def _collect_alias_review(ingredients_text: str, review_counter: dict[str, dict[str, Any]]) -> None:
    normalized = ingredients_text.lower().replace("ё", "е")
    tokens = [token.strip() for token in re.split(r"[,;/]", normalized) if token.strip()]
    for token in tokens:
        token_clean = re.sub(r"\s+", " ", re.sub(r"[^a-zа-я0-9\s-]", " ", token)).strip()
        if not token_clean:
            continue
        canonical = ALIASES.get(token_clean)
        status = "mapped" if canonical else "needs_review"
        key = f"{token_clean}|{status}|{canonical or ''}"
        if key not in review_counter:
            review_counter[key] = {
                "alias": token_clean,
                "count": 0,
                "canonical": canonical or "",
                "status": status,
            }
        review_counter[key]["count"] += 1


def _write_dish_audit(
    audit_dir: Path,
    loaded: dict[str, Any],
    dishes: list[dict[str, Any]],
    cache_rows: list[dict[str, Any]],
    evidence_rows: list[list[Any]],
    missing_primary_rows: list[list[Any]],
) -> None:
    summary_path = audit_dir / "dish_stage_a_summary.csv"
    with summary_path.open("w", encoding="utf-8", newline="") as file_obj:
        writer = csv.writer(file_obj)
        writer.writerow(
            [
                "parsed_rows",
                "parsed_dishes",
                "skipped_rows",
                "skipped_by_reason",
                "unique_dish_ids",
                "cache_rows",
                "nutrition_parse_success_rate",
            ]
        )
        nutrition_rate = 0.0
        if len(dishes) > 0:
            nutrition_rate = loaded["nutrition_success"] / len(dishes)
        writer.writerow(
            [
                loaded["parsed_rows"],
                len(dishes),
                loaded["skipped_rows"],
                json.dumps(loaded["skipped_by_reason"], ensure_ascii=False),
                len({str(d["dish_id"]) for d in dishes}),
                len(cache_rows),
                round(nutrition_rate, 4),
            ]
        )

    skipped_path = audit_dir / "dish_stage_a_skipped_rows.csv"
    with skipped_path.open("w", encoding="utf-8", newline="") as file_obj:
        writer = csv.writer(file_obj)
        writer.writerow(["row_index", "reason", "dish_id", "dish_name"])
        for row in loaded["skipped_rows_detail"]:
            writer.writerow([row["row_index"], row["reason"], row["dish_id"], row["dish_name"]])

    nutrition_errors_path = audit_dir / "nutrition_parse_errors.csv"
    with nutrition_errors_path.open("w", encoding="utf-8", newline="") as file_obj:
        writer = csv.writer(file_obj)
        writer.writerow(["row_index", "dish_id", "dish_name", "raw_kcal"])
        for row in loaded["nutrition_parse_errors"]:
            writer.writerow([row["row_index"], row["dish_id"], row["dish_name"], row["raw_kcal"]])

    review_path = audit_dir / "ingredient_alias_review.csv"
    with review_path.open("w", encoding="utf-8", newline="") as file_obj:
        writer = csv.writer(file_obj)
        writer.writerow(["alias", "count", "canonical", "status"])
        for row in sorted(loaded["ingredient_alias_review"], key=lambda x: (-x["count"], x["alias"])):
            writer.writerow([row["alias"], row["count"], row["canonical"], row["status"]])

    evidence_path = audit_dir / "stage_a_feature_evidence_sample.csv"
    with evidence_path.open("w", encoding="utf-8", newline="") as file_obj:
        writer = csv.writer(file_obj)
        writer.writerow(["dish_id", "dish_name", "axis", "feature_key", "evidence", "confidence_or_value", "source"])
        for row in evidence_rows[:300]:
            writer.writerow(row)

    missing_primary_path = audit_dir / "stage_a_missing_primary_ingredient.csv"
    with missing_primary_path.open("w", encoding="utf-8", newline="") as file_obj:
        writer = csv.writer(file_obj)
        writer.writerow(["dish_id", "dish_name", "ingredient_key", "reason"])
        for row in missing_primary_rows:
            writer.writerow(row)


def _write_false_positive_watchlist(audit_dir: Path, cache_rows: list[dict[str, Any]]) -> None:
    watchlist_path = audit_dir / "stage_a_false_positive_watchlist.csv"
    candidates = [
        row
        for row in cache_rows
        if (
            "koreanchick" in _normalize(row.get("meta", {}).get("restaurant", ""))
            or "кореан" in _normalize(row.get("meta", {}).get("restaurant", ""))
        )
        and (
            any(
                token
                in (
                    _normalize(row.get("meta", {}).get("name", ""))
                    + " "
                    + _normalize(row.get("meta", {}).get("category", ""))
                )
                for token in ("вок", "рис", "wok", "rice")
            )
            or any(token in row.get("ingredient", {}) for token in ("rice", "noodles"))
        )
    ]
    sample = candidates[:20]
    asian_evidence_hits = 0
    asian_predicted = 0
    salad_fp = 0
    with watchlist_path.open("w", encoding="utf-8", newline="") as file_obj:
        writer = csv.writer(file_obj)
        writer.writerow(
            [
                "dish_id",
                "dish_name",
                "restaurant",
                "explicit_asian_evidence",
                "asian_tag",
                "salad_tag",
                "result",
            ]
        )
        for row in sample:
            name = _normalize(row.get("meta", {}).get("name", ""))
            ingredients = _normalize(" ".join(row.get("ingredient", {}).keys()))
            explicit = any(keyword in f"{name} {ingredients}" for keyword in ASIAN_EXPLICIT_KEYWORDS)
            asian_tag = "asian" in row.get("cuisine", {})
            salad_tag = "salad" in row.get("format_texture", {})
            if explicit:
                asian_evidence_hits += 1
            if asian_tag:
                asian_predicted += 1
            if salad_tag:
                salad_fp += 1
            result = "ok"
            if explicit and not asian_tag:
                result = "asian_missed"
            if salad_tag:
                result = "salad_false_positive"
            writer.writerow(
                [
                    row.get("dish_id"),
                    row.get("meta", {}).get("name", ""),
                    row.get("meta", {}).get("restaurant", ""),
                    explicit,
                    asian_tag,
                    salad_tag,
                    result,
                ]
            )
        precision = 0.0
        if asian_predicted > 0:
            true_positive = sum(
                1
                for row in sample
                if ("asian" in row.get("cuisine", {}))
                and any(
                    keyword in _normalize(row.get("meta", {}).get("name", "") + " " + " ".join(row.get("ingredient", {}).keys()))
                    for keyword in ASIAN_EXPLICIT_KEYWORDS
                )
            )
            precision = true_positive / asian_predicted
        writer.writerow([])
        writer.writerow(["sample_size", len(sample)])
        writer.writerow(["asian_precision", round(precision, 4)])
        writer.writerow(["salad_false_positives", salad_fp])
        gate_pass = len(sample) >= 20 and precision >= 0.9 and salad_fp == 0
        writer.writerow(["gate_pass", gate_pass])
